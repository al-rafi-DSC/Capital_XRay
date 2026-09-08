"""Capital X-Ray - build the province x sector x year dataset.

Reads Banca d'Italia STAFINRA cubes (loans, deposits) and Istat Frame SBS
territorial tables (turnover, business counts), joins them at
province x macro-sector x year, computes the Capital Gap Score and writes
web/data.json for the frontend.

Run:  python pipeline/build_dataset.py
"""
import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = ROOT / "web" / "data.json"

# Banca d'Italia NUTS province codes -> display name
PROVINCES = {
    "ITC4C": "Milano",
    "ITC47": "Brescia",
    "ITC46": "Bergamo",
    "ITC4D": "Monza e della Brianza",
    "ITC41": "Varese",
}
# Banca d'Italia ATECO_CTP codes -> macro-sector
BDI_SECTORS = {"1005001": "Industria", "1005003": "Servizi", "F": "Costruzioni"}
# Istat ATECO section -> macro-sector (crosswalk)
ISTAT_INDUSTRIA = {"B_e_C", "D", "E"}
ISTAT_COSTRUZIONI = {"F"}

YEARS = [str(y) for y in range(2015, 2021)]

# Institutional-sector filters
LOANS_SET_CTP = "SBI25"     # societa non finanziarie e famiglie produttrici
DEPOSITS_SET_CTP = "SBI42"  # totale residenti al netto delle IFM

CUBE_LOANS = "TFR20232"
CUBE_DEPOSITS = "TFR20267"


def read_cube(cube):
    """Pull one STAFINRA cube out of the nested zips as a DataFrame."""
    outer = zipfile.ZipFile(DATA / "raw" / "banca_finanziamenti_raccolta.zip")
    data_zip = zipfile.ZipFile(io.BytesIO(outer.read("STAFINRA_DATA.zip")))
    entry = next(n for n in data_zip.namelist() if n.split("/")[-1].startswith(cube))
    cube_zip = zipfile.ZipFile(io.BytesIO(data_zip.read(entry)))
    csv_name = cube_zip.namelist()[0]
    return pd.read_csv(
        io.BytesIO(cube_zip.read(csv_name)), sep=";", encoding="latin-1", dtype=str
    )


def load_loans():
    df = read_cube(CUBE_LOANS)
    df["value"] = pd.to_numeric(df.VALORE, errors="coerce")
    df = df[
        df.LOC_CTP.isin(PROVINCES)
        & df.ATECO_CTP.isin(BDI_SECTORS)
        & (df.SET_CTP == LOANS_SET_CTP)
        & df.DATA_OSS.str.endswith("12-31")
    ].copy()
    df["province"] = df.LOC_CTP.map(PROVINCES)
    df["sector"] = df.ATECO_CTP.map(BDI_SECTORS)
    df["year"] = df.DATA_OSS.str[:4]
    return (
        df.groupby(["province", "sector", "year"], as_index=False)["value"]
        .sum()
        .rename(columns={"value": "loans_kEUR"})
    )


def load_deposits():
    """Deposits carry no ATECO dimension - they stay province-level."""
    df = read_cube(CUBE_DEPOSITS)
    df["value"] = pd.to_numeric(df.VALORE, errors="coerce")
    df = df[
        df.LOC_CTP.isin(PROVINCES)
        & (df.SET_CTP == DEPOSITS_SET_CTP)
        & df.DATA_OSS.str.endswith("12-31")
    ].copy()
    df["province"] = df.LOC_CTP.map(PROVINCES)
    df["year"] = df.DATA_OSS.str[:4]
    return (
        df.groupby(["province", "year"], as_index=False)["value"]
        .sum()
        .rename(columns={"value": "deposits_kEUR"})
    )


def load_istat():
    """One workbook per year, one sheet per ATECO section."""
    rows = []
    for year in YEARS:
        path = (
            DATA
            / "extracted/istat_frame_sbs_2023/Tavole/Province 2015-2020"
            / f"Frame territoriale - Province - Anno {year}.xlsx"
        )
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in wb.sheetnames:
            match = re.search(r"Province.*Sez_+(.+)$", sheet)
            if not match:
                continue
            section = match.group(1)
            if section in ISTAT_INDUSTRIA:
                sector = "Industria"
            elif section in ISTAT_COSTRUZIONI:
                sector = "Costruzioni"
            else:
                sector = "Servizi"
            for r in wb[sheet].iter_rows(values_only=True):
                # col 3 = province, 4 = unita locali, 5 = addetti, 10 = fatturato
                if (
                    len(r) > 10
                    and r[3] in PROVINCES.values()
                    and isinstance(r[4], (int, float))
                ):
                    rows.append(
                        {
                            "province": r[3],
                            "sector": sector,
                            "year": year,
                            "businesses": r[4],
                            "employees": r[5],
                            "turnover_kEUR": r[10],
                        }
                    )
        wb.close()
    return (
        pd.DataFrame(rows)
        .groupby(["province", "sector", "year"], as_index=False)[
            ["businesses", "employees", "turnover_kEUR"]
        ]
        .sum()
    )


def cagr(first, last, periods):
    if not first or not last or first <= 0 or periods <= 0:
        return 0.0
    return ((last / first) ** (1 / periods) - 1) * 100


def scale_0_100(series):
    lo, hi = series.min(), series.max()
    if hi == lo:
        return pd.Series([50.0] * len(series), index=series.index)
    return (series - lo) / (hi - lo) * 100


def build():
    loans, deposits, istat = load_loans(), load_deposits(), load_istat()
    df = istat.merge(loans, on=["province", "sector", "year"], how="inner").merge(
        deposits, on=["province", "year"], how="inner"
    )
    df = df[df.year.isin(YEARS)].sort_values(["province", "sector", "year"])

    first_year, last_year = YEARS[0], YEARS[-1]
    span = len(YEARS) - 1
    records = []
    for (province, sector), g in df.groupby(["province", "sector"]):
        g = g.set_index("year")
        if first_year not in g.index or last_year not in g.index:
            continue
        a, b = g.loc[first_year], g.loc[last_year]
        records.append(
            {
                "province": province,
                "sector": sector,
                "turnover_growth": cagr(a.turnover_kEUR, b.turnover_kEUR, span),
                "business_growth": cagr(a.businesses, b.businesses, span),
                "employee_growth": cagr(a.employees, b.employees, span),
                "credit_growth": cagr(a.loans_kEUR, b.loans_kEUR, span),
                "deposit_growth": cagr(a.deposits_kEUR, b.deposits_kEUR, span),
                "turnover_kEUR": float(b.turnover_kEUR),
                "businesses": int(b.businesses),
                "loans_kEUR": float(b.loans_kEUR),
                "deposits_kEUR": float(b.deposits_kEUR),
                "credit_intensity": float(b.loans_kEUR) / float(b.turnover_kEUR) * 100,
                "series": [
                    {
                        "year": y,
                        "turnover_kEUR": float(g.loc[y].turnover_kEUR),
                        "loans_kEUR": float(g.loc[y].loans_kEUR),
                        "businesses": int(g.loc[y].businesses),
                    }
                    for y in YEARS
                    if y in g.index
                ],
            }
        )

    out = pd.DataFrame(records)

    # --- sub-scores, each normalised 0-100 across the 15-cell grid ---
    # Liquidity = deposit coverage of lending. Deposits are province-level and
    # loans are sector-level, so the ratio still varies by sector.
    out["liquidity_ratio"] = out.deposits_kEUR / out.loans_kEUR

    out["s_turnover_growth"] = scale_0_100(out.turnover_growth)
    out["s_business_growth"] = scale_0_100(out.business_growth)
    out["s_sector_performance"] = scale_0_100(out.employee_growth)
    out["s_credit_growth"] = scale_0_100(out.credit_growth)
    out["s_deposit_strength"] = scale_0_100(out.deposit_growth)
    out["s_liquidity"] = scale_0_100(out.liquidity_ratio)

    # PRD section 5 weights
    out["economic_momentum"] = (
        0.4 * out.s_turnover_growth
        + 0.3 * out.s_business_growth
        + 0.3 * out.s_sector_performance
    )
    out["financing_support"] = (
        0.5 * out.s_credit_growth
        + 0.3 * out.s_deposit_strength
        + 0.2 * out.s_liquidity
    )
    # Frontend sign convention: NEGATIVE = underfinanced (financing deficit).
    # This is the inverse of the PRD's prose, and matches app/page.tsx.
    out["capital_gap"] = out.financing_support - out.economic_momentum

    out = out.sort_values("capital_gap")

    def slug(text):
        text = text.lower().replace("'", "").replace(" ", "-")
        return re.sub(r"[^a-z0-9-]", "", text)

    rows = []
    for r in out.itertuples():
        rows.append(
            {
                "id": f"{slug(r.province)}-{slug(r.sector)}",
                "province": r.province,
                "sector": r.sector,
                "economicMomentum": round(r.economic_momentum),
                "financingSupport": round(r.financing_support),
                "capitalGapScore": round(r.capital_gap),
                "subScores": {
                    "turnoverGrowth": round(r.s_turnover_growth),
                    "businessGrowth": round(r.s_business_growth),
                    "sectorPerformance": round(r.s_sector_performance),
                    "creditGrowth": round(r.s_credit_growth),
                    "depositStrength": round(r.s_deposit_strength),
                    "liquidity": round(r.s_liquidity),
                },
                # raw values kept for the AI prompt and tooltips
                "raw": {
                    "turnoverGrowthPct": round(r.turnover_growth, 2),
                    "businessGrowthPct": round(r.business_growth, 2),
                    "employeeGrowthPct": round(r.employee_growth, 2),
                    "creditGrowthPct": round(r.credit_growth, 2),
                    "depositGrowthPct": round(r.deposit_growth, 2),
                    "turnoverKEUR": r.turnover_kEUR,
                    "loansKEUR": r.loans_kEUR,
                    "depositsKEUR": r.deposits_kEUR,
                    "businesses": r.businesses,
                    "creditIntensityPct": round(r.credit_intensity, 2),
                },
                "series": r.series,
            }
        )

    meta = {
        "years": YEARS,
        "provinces": sorted(out.province.unique().tolist()),
        "sectors": sorted(out.sector.unique().tolist()),
        "sources": {
            "credit": f"Banca d'Italia STAFINRA {CUBE_LOANS} / {CUBE_DEPOSITS}",
            "economy": "Istat Frame SBS territoriale, province 2015-2020",
        },
        "notes": "Negative capitalGapScore = underfinanced. Deposits are province-level.",
    }

    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(
        json.dumps({"meta": meta, "cells": rows}, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    # TypeScript module consumed directly by app/page.tsx (no fetch, no CORS)
    ts_path = ROOT / "capital-x-ray" / "lib" / "capital-data.ts"
    if ts_path.parent.exists():
        banner = (
            "// GENERATED by pipeline/build_dataset.py - do not edit by hand.\n"
            f"// Sources: {meta['sources']['credit']}; {meta['sources']['economy']}\n"
            "// Negative capitalGapScore = underfinanced (financing deficit).\n\n"
        )
        body = json.dumps(rows, indent=2, ensure_ascii=False)
        ts_path.write_text(
            f"{banner}export const CAPITAL_META = "
            f"{json.dumps(meta, indent=2, ensure_ascii=False)};\n\n"
            f"export const CAPITAL_DATA = {body};\n",
            encoding="utf-8",
        )
    return out


if __name__ == "__main__":
    result = build()
    print(f"wrote {OUT}  ({len(result)} cells)")
    cols = [
        "province",
        "sector",
        "economic_momentum",
        "financing_support",
        "capital_gap",
    ]
    print(result[cols].to_string(index=False, float_format=lambda v: f"{v:7.1f}"))
    print("\n(negative capitalGapScore = underfinanced, matching the frontend)")
